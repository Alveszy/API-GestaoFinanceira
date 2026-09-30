from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from pathlib import Path
import re

root = Path(r'C:\Users\mateu\APIgestao')
source = root / 'docs' / 'arquitetura-e-processo.md'
out = root / 'docs' / 'Documentacao_Arquitetural_API_Gestao_Financeira_Revisada.docx'
image_map = {
    'C4 — Containers': root / 'docs' / 'c4_containers.png',
    'C4 — Componentes da API': root / 'docs' / 'c4_components.png',
    'UML — Diagrama de classes do modelo': root / 'docs' / 'uml_classes.png',
}

# Context diagram rendered as a clean raster figure for the Word version.
context_path = root / 'docs' / 'c4_context.png'
canvas = Image.new('RGB', (1600, 560), '#f8fafc')
draw = ImageDraw.Draw(canvas)
font_path = r'C:\Windows\Fonts\arial.ttf'
bold_path = r'C:\Windows\Fonts\arialbd.ttf'
f_title = ImageFont.truetype(bold_path, 27)
f_body = ImageFont.truetype(font_path, 22)
f_label = ImageFont.truetype(bold_path, 19)
navy, teal, pale, ink = '#173653', '#007f8b', '#eaf2f7', '#283644'
def card(box, title, lines):
    x1,y1,x2,y2=box
    draw.rounded_rectangle(box, radius=20, fill='white', outline=teal, width=4)
    draw.rounded_rectangle((x1,y1,x2,y1+65), radius=20, fill=navy)
    draw.rectangle((x1,y1+42,x2,y1+65), fill=navy)
    draw.text((x1+22,y1+14),title,font=f_title,fill='white')
    for j,line in enumerate(lines): draw.text((x1+22,y1+88+j*39),line,font=f_body,fill=ink)
card((65,190,430,390),'Pessoa usuária',['Cliente ou administrador','Navegador / cliente HTTP'])
card((600,150,1090,430),'API de Gestão Financeira',['Aplicação REST em NestJS','Regras financeiras e autenticação','Swagger UI em /docs'])
card((1200,190,1535,390),'Swagger UI',['Interface interativa em /docs','Apresenta o contrato OpenAPI'])
draw.line((430,275,600,275),fill=teal,width=6); draw.polygon([(600,275),(575,260),(575,290)],fill=teal)
draw.text((445,232),'HTTPS / JSON / JWT',font=f_label,fill='#526b7d')
draw.line((1090,320,1200,320),fill=teal,width=5); draw.polygon([(1200,320),(1175,305),(1175,335)],fill=teal)
draw.text((1110,275),'OpenAPI',font=f_label,fill='#526b7d')
canvas.save(context_path)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_text(cell, text, header=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.strip())
    run.font.name = 'Aptos'; run.font.size = Pt(9)
    if header:
        run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header: shade(cell, '173653')

def add_rich(p, text):
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0,108,125)
        else:
            p.add_run(part)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.68); sec.bottom_margin = Inches(.68); sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Aptos'; normal.font.size = Pt(10); normal.font.color.rgb = RGBColor(40,54,68)
normal.paragraph_format.space_after = Pt(3); normal.paragraph_format.line_spacing = 1.0
for name, size, color in [('Title', 26, '173653'), ('Heading 1', 17, '173653'), ('Heading 2', 12, '007F8B'), ('Heading 3', 10, '426278')]:
    st=styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(12 if name!='Title' else 0); st.paragraph_format.space_after=Pt(5)
title_ppr = styles['Title']._element.pPr
if title_ppr is not None:
    title_border = title_ppr.find(qn('w:pBdr'))
    if title_border is not None: title_ppr.remove(title_border)

# Header and page-number footer
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rh = header.add_run('API DE GESTÃO FINANCEIRA  |  DOCUMENTAÇÃO ARQUITETURAL')
rh.font.name='Aptos'; rh.font.size=Pt(8); rh.font.color.rgb=RGBColor(95,115,130)
footer = sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
rf=footer.add_run('Documentação arquitetural  •  '); rf.font.size=Pt(8); rf.font.color.rgb=RGBColor(95,115,130)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(82); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
r=p.add_run('ARQUITETURA DE SOFTWARE'); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor(0,127,139)
p=doc.add_paragraph(style='Title'); p.paragraph_format.space_before=Pt(9); p.add_run('API de Gestão Financeira')
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(8)
r=p.add_run('Processo de criação, modelo C4 e modelagem UML'); r.font.size=Pt(15); r.font.color.rgb=RGBColor(66,98,120)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(24)
p.add_run('Documento técnico do projeto acadêmico').font.size=Pt(11)
p=doc.add_paragraph(); p.add_run('Arquitetura em camadas  •  NestJS  •  PostgreSQL  •  Redis').font.color.rgb=RGBColor(0,127,139)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(120)
p.add_run('29 de setembro de 2026').font.size=Pt(10)
p=doc.add_paragraph(); p.add_run('Escopo: processo de criação, diagramas C4, classes UML e requisitos não funcionais implementados.').font.size=Pt(9)
doc.add_page_break()

lines=source.read_text(encoding='utf-8').splitlines()
in_code=False; code_title=''; table_rows=[]; i=0; skip_first=True
uml_summary_next=False
while i < len(lines):
    line=lines[i]
    if line.startswith('# '):
        i+=1; continue
    if line.startswith('```'):
        if in_code:
            title_key = code_title.lower()
            img = image_map.get(code_title)
            if img is None and 'diagrama de containers' in title_key:
                img = root / 'docs' / 'c4_containers.png'
            elif img is None and 'componentes' in title_key:
                img = root / 'docs' / 'c4_components.png'
            elif img is None and 'classes do modelo' in title_key:
                img = root / 'docs' / 'uml_classes.png'
            elif img is None and 'diagrama de contexto' in title_key:
                img = context_path
            if img and img.exists():
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(4)
                if 'classes do modelo' in title_key: p.paragraph_format.keep_together=True
                p.add_run().add_picture(str(img), width=Inches(6.65))
                cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(7)
                note = 'O campo Account.type é armazenado como texto no código atual.' if 'classes do modelo' in title_key else 'Representação baseada nos módulos e serviços implementados.'
                rr=cap.add_run(code_title + ' — ' + note); rr.italic=True; rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(95,115,130)
            in_code=False; code_title=''
        else:
            in_code=True
            prev = next((x[2:].strip() for x in reversed(lines[:i]) if x.startswith('## ')), '')
            code_title=prev
        i+=1; continue
    if in_code:
        i+=1; continue
    if line.startswith('|'):
        cells=[c.strip() for c in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r'[-: ]+', c or '-') for c in cells):
            i+=1; continue
        table_rows.append(cells); i+=1
        while i < len(lines) and lines[i].startswith('|'):
            cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', c or '-') for c in cells): table_rows.append(cells)
            i+=1
        if table_rows:
            cols=max(map(len,table_rows)); t=doc.add_table(rows=0, cols=cols); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
            for ri,row in enumerate(table_rows):
                cells=t.add_row().cells
                if ri == 0:
                    trPr = t.rows[-1]._tr.get_or_add_trPr()
                    repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'), 'true'); trPr.append(repeat)
                for ci in range(cols): set_cell_text(cells[ci], row[ci] if ci<len(row) else '', header=(ri==0))
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
            table_rows=[]
        continue
    if not line.strip(): i+=1; continue
    if line.startswith('### '): doc.add_paragraph(line[4:].strip(), style='Heading 3')
    elif line.startswith('## '):
        heading = doc.add_paragraph(line[3:].strip(), style='Heading 1')
        if 'Diagrama de classes do modelo' in line: heading.paragraph_format.keep_with_next=True; uml_summary_next=True
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); add_rich(p,line[2:].strip())
    elif re.match(r'^\d+\.\s',line):
        p=doc.add_paragraph(); add_rich(p,line.strip())
    else:
        p=doc.add_paragraph(); add_rich(p,line.strip())
        if uml_summary_next:
            p.paragraph_format.keep_with_next=True; uml_summary_next=False
    i+=1

# Diagram-specific notes clarify the UML account type is stored as text in the implementation.
# Document metadata
doc.core_properties.title='Documentação Arquitetural da API de Gestão Financeira'
doc.core_properties.subject='Processo de criação, arquitetura C4, UML e requisitos não funcionais'
doc.core_properties.author='Projeto API de Gestão Financeira'
doc.save(out)
print(out)
