from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)
NAVY = "17324D"
TEAL = "087E8B"
PALE = "EAF2F7"
MID = "536879"
INK = "263746"
WHITE = "FFFFFF"
FONT = "C:/Windows/Fonts/arial.ttf"
FONT_B = "C:/Windows/Fonts/arialbd.ttf"


def rgb(hex_color):
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))


def fnt(size, bold=False):
    return ImageFont.truetype(FONT_B if bold else FONT, size)


def diagram_canvas(size, title):
    im = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(im)
    d.text((50, 30), title, fill=rgb(NAVY), font=fnt(36, True))
    return im, d


def card(d, box, title, lines, fill=PALE, title_color=NAVY, body_color=INK, title_size=27, body_size=21):
    x1, y1, x2, y2 = box
    fill_color = rgb(fill) if isinstance(fill, str) else fill
    d.rounded_rectangle(box, radius=18, fill=fill_color, outline=rgb(TEAL), width=3)
    d.rounded_rectangle((x1, y1, x2, y1+52), radius=16, fill=rgb(NAVY))
    d.rectangle((x1, y1+32, x2, y1+52), fill=rgb(NAVY))
    d.text((x1+18, y1+9), title, fill=rgb(WHITE) if fill_color == rgb(PALE) else rgb(title_color), font=fnt(title_size, True))
    yy = y1 + 68
    for line in lines:
        d.text((x1+18, yy), line, fill=rgb(body_color), font=fnt(body_size))
        yy += body_size + 10


def arrow(d, points, color=TEAL, width=5, label=None, label_xy=None):
    d.line(points, fill=rgb(color), width=width, joint="curve")
    x0, y0 = points[-2]
    x1, y1 = points[-1]
    import math
    ang = math.atan2(y1-y0, x1-x0)
    length = 18
    p1 = (x1-length*math.cos(ang-0.5), y1-length*math.sin(ang-0.5))
    p2 = (x1-length*math.cos(ang+0.5), y1-length*math.sin(ang+0.5))
    d.polygon([(x1, y1), p1, p2], fill=rgb(color))
    if label and label_xy:
        d.rounded_rectangle((label_xy[0]-8,label_xy[1]-3,label_xy[0]+len(label)*12,label_xy[1]+28), radius=5, fill="white")
        d.text(label_xy, label, fill=rgb(MID), font=fnt(18, True))


def save_c4_container(path):
    im, d = diagram_canvas((1800, 760), "C4 | Diagrama de Contêineres")
    card(d, (65, 300, 365, 465), "Pessoa usuária", ["Cliente / administrador", "Usa navegador ou cliente HTTP"], fill="FFFFFF")
    card(d, (520, 205, 1040, 555), "API de Gestão Financeira", ["NestJS + TypeScript", "REST /api/v1 e Swagger /docs", "Apresentação, aplicação e infraestrutura", "JWT, validação e regras financeiras"])
    card(d, (1245, 115, 1705, 325), "PostgreSQL 16", ["Banco relacional", "Usuários, contas, categorias,", "transações e orçamentos"], fill="FFFFFF")
    card(d, (1245, 440, 1705, 650), "Redis 7", ["Cache operacional", "Chaves e respostas de", "idempotência (24 horas)"], fill="FFFFFF")
    arrow(d, [(365, 380), (520, 380)], label="REST + JWT", label_xy=(390, 335))
    arrow(d, [(1040, 300), (1140, 300), (1140, 220), (1245, 220)], label="SQL / TypeORM", label_xy=(1060, 250))
    arrow(d, [(1040, 465), (1140, 465), (1140, 545), (1245, 545)], label="Redis", label_xy=(1060, 480))
    im.save(path)


def save_c4_components(path):
    im, d = diagram_canvas((1800, 1060), "C4 | Diagrama de Componentes da API")
    d.rounded_rectangle((55, 100, 1745, 995), radius=24, outline=rgb(NAVY), width=5, fill="#FBFCFE")
    d.text((82, 115), "Contêiner: API NestJS", fill=rgb(NAVY), font=fnt(28, True))
    card(d, (115, 200, 590, 385), "Apresentação", ["Controllers HTTP", "DTOs e validação", "Swagger / OpenAPI"])
    card(d, (690, 180, 1120, 355), "Aplicação Auth", ["Cadastro e login", "Hash bcrypt e JWT", "Autorização por papéis"])
    card(d, (690, 425, 1120, 650), "Aplicação Financeira", ["Contas e categorias", "Transações e orçamentos", "Isolamento por usuário"])
    card(d, (115, 545, 590, 785), "Serviços transversais", ["Problem Details (RFC 7807)", "Correlation-ID e logs JSON", "Idempotência com Redis", "Rate limiting"])
    card(d, (1240, 310, 1645, 650), "Infraestrutura", ["TypeORM / PostgreSQL", "Persistência relacional", "ioredis / Redis", "Respostas idempotentes"])
    arrow(d, [(590, 285), (690, 285)])
    arrow(d, [(590, 320), (630, 320), (630, 505), (690, 505)])
    arrow(d, [(590, 675), (640, 675), (640, 820), (1180, 820), (1180, 620), (1240, 620)])
    arrow(d, [(1120, 270), (1180, 270), (1180, 420), (1240, 420)])
    arrow(d, [(1120, 535), (1240, 535)])
    d.text((115, 850), "A apresentação encaminha requisições aos componentes de aplicação; estes usam a infraestrutura para persistir", fill=rgb(MID), font=fnt(22))
    d.text((115, 885), "dados financeiros ou consultar idempotência. Categorias e orçamentos acessam TypeORM diretamente.", fill=rgb(MID), font=fnt(22))
    im.save(path)


def save_uml(path):
    im, d = diagram_canvas((1900, 1280), "UML | Diagrama de Classes do Domínio Financeiro")
    boxes = {
        "user": (700, 115, 1200, 390),
        "account": (70, 570, 480, 895),
        "category": (515, 570, 925, 895),
        "transaction": (960, 570, 1420, 955),
        "budget": (1450, 570, 1830, 895),
    }
    def umlbox(key, name, attributes):
        x1,y1,x2,y2 = boxes[key]
        d.rectangle((x1,y1,x2,y2), fill="white", outline=rgb(NAVY), width=3)
        d.rectangle((x1,y1,x2,y1+55), fill=rgb(NAVY))
        d.text((x1+15,y1+11), name, fill=rgb(WHITE), font=fnt(27, True))
        yy=y1+74
        for a in attributes:
            d.text((x1+16, yy), a, fill=rgb(INK), font=fnt(19))
            yy += 35
        d.line((x1,y1+55,x2,y1+55), fill=rgb(TEAL), width=2)
    # Draw associations behind the class boxes.
    arrow(d, [(820,390),(820,475),(275,475),(275,570)], color=MID, width=3)
    arrow(d, [(900,390),(900,490),(720,490),(720,570)], color=MID, width=3)
    arrow(d, [(1000,390),(1000,510),(1190,510),(1190,570)], color=MID, width=3)
    arrow(d, [(1100,390),(1100,485),(1640,485),(1640,570)], color=MID, width=3)
    arrow(d, [(480,720),(960,720)], color=TEAL, width=3)
    arrow(d, [(925,810),(960,810)], color=TEAL, width=3)
    arrow(d, [(925,855),(1450,855)], color=TEAL, width=3)
    umlbox("user", "User", ["+ id: UUID", "+ name: String", "+ email: String (unique)", "- passwordHash: String", "+ role: admin | cliente", "+ createdAt, updatedAt: Date"])
    umlbox("account", "Account", ["+ id: UUID", "+ name: String", "+ type: AccountType", "+ currency: CHAR(3)", "+ openingBalance: Decimal", "+ active: Boolean"])
    umlbox("category", "Category", ["+ id: UUID", "+ name: String", "+ kind: income | expense", "+ color: String?", "+ createdAt: Date"])
    umlbox("transaction", "Transaction", ["+ id: UUID", "+ kind: income | expense", "+ amount: Decimal", "+ description: String", "+ occurredAt: Date", "+ status: TransactionStatus"])
    umlbox("budget", "Budget", ["+ id: UUID", "+ month: YYYY-MM", "+ limitAmount: Decimal", "+ name: String?", "+ createdAt, updatedAt: Date"])
    d.text((80, 1010), "Associações", fill=rgb(NAVY), font=fnt(23, True))
    d.text((80, 1050), "User 1 -> 0..* Account, Category, Transaction, Budget | Account 1 -> 0..* Transaction", fill=rgb(INK), font=fnt(20))
    d.text((80, 1090), "Category 0..1 -> 0..* Transaction e Budget | Valores monetários persistidos como NUMERIC(14,2)", fill=rgb(INK), font=fnt(20))
    d.text((80, 1145), "FKs para usuário usam exclusão em cascata; transações preservam integridade da conta (RESTRICT).", fill=rgb(MID), font=fnt(18))
    im.save(path)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=INK, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, head in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], head, True, WHITE, 9)
        shade(table.rows[0].cells[i], NAVY)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, value in enumerate(row):
            set_cell_text(cells[ci], str(value), False, INK, 8.5)
            if ri % 2 == 0:
                shade(cells[ci], "F2F6F9")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("API Gestão Financeira  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(.65)
    sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.72)
    sec.right_margin = Inches(.72)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color in [("Title", 30, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 11, NAVY)]:
        s = styles[name]
        s.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        s.paragraph_format.space_after = Pt(5)
    title_ppr = styles["Title"]._element.pPr
    if title_ppr is not None:
        title_border = title_ppr.find(qn("w:pBdr"))
        if title_border is not None:
            title_ppr.remove(title_border)
    add_page_field(sec.footer.paragraphs[0])


def heading(doc, title, level=1):
    p = doc.add_paragraph(title, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_image(doc, path, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    return p


def caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    for r in p.runs:
        r.italic = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(MID)


def build():
    uml = OUT / "uml_classes.png"
    c4c = OUT / "c4_containers.png"
    c4p = OUT / "c4_components.png"
    save_uml(uml)
    save_c4_container(c4c)
    save_c4_components(c4p)

    doc = Document()
    setup(doc)
    doc.core_properties.title = "Documentação Técnica da API de Gestão Financeira"
    doc.core_properties.subject = "Arquitetura em camadas, modelo UML e diagramas C4"
    doc.core_properties.keywords = "NestJS, PostgreSQL, Redis, UML, C4, API"

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("DOCUMENTAÇÃO TÉCNICA")
    r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(TEAL)
    title = doc.add_paragraph("API de Gestão Financeira", style="Title")
    title.paragraph_format.space_after = Pt(12)
    sub = doc.add_paragraph("Arquitetura em camadas, modelo de dados UML e diagramas C4")
    sub.paragraph_format.space_after = Pt(32)
    for run in sub.runs:
        run.font.size = Pt(15); run.font.color.rgb = RGBColor.from_string(MID)
    add_table(doc, ["Resumo do projeto", "Descrição"], [
        ("Objetivo", "Disponibilizar uma API REST para gerenciar contas, categorias, receitas, despesas e orçamentos pessoais."),
        ("Stack", "NestJS, TypeScript, TypeORM, PostgreSQL 16, Redis 7 e Docker Compose."),
        ("Contrato", "OpenAPI/Swagger em /docs; endpoints versionados sob /api/v1."),
        ("Arquitetura", "Arquitetura em camadas, organizada em módulos por domínio."),
    ], [1.25, 5.15])
    p = doc.add_paragraph("Documento de arquitetura e implementação")
    p.paragraph_format.space_before = Pt(35)
    p.runs[0].font.size = Pt(11); p.runs[0].font.color.rgb = RGBColor.from_string(MID)

    doc.add_page_break()
    heading(doc, "1. Visão geral", 1)
    doc.add_paragraph("A API de Gestão Financeira oferece operações autenticadas para organizar contas e registrar movimentações financeiras. Cada cliente acessa apenas os próprios dados. Administradores podem consultar a lista de usuários, e o provisionamento administrativo é feito por configuração de ambiente.")
    heading(doc, "Em termos simples", 2)
    doc.add_paragraph("Pense na API como o serviço que fica nos bastidores de um aplicativo financeiro. Ela recebe pedidos como “criar uma conta” ou “registrar uma despesa”, confere quem está fazendo o pedido e guarda as informações nos bancos de dados. Este projeto implementa esse serviço e sua documentação; não inclui uma tela de aplicativo para o cliente final.")
    heading(doc, "Objetivos", 2)
    for item in [
        "Registrar receitas e despesas vinculadas a contas e categorias.",
        "Definir limites mensais de orçamento, opcionais por categoria.",
        "Documentar o contrato HTTP e os modelos usados nas requisições.",
        "Aplicar autenticação, autorização por papel, idempotência, limitação de requisições, erros padronizados e correlação de logs.",
    ]: add_bullet(doc, item)
    heading(doc, "Escopo funcional", 2)
    add_table(doc, ["Módulo", "Responsabilidade"], [
        ("Autenticação", "Cadastro de cliente, login JWT e autorização admin/cliente."),
        ("Contas", "Criar, listar e remover contas financeiras do usuário autenticado."),
        ("Categorias", "Criar e listar categorias de receita ou despesa."),
        ("Transações", "Criar, listar com filtros de data e remover movimentações."),
        ("Orçamentos", "Criar e listar limites mensais gerais ou associados a categorias."),
        ("Saúde", "Verificar disponibilidade da API e conectividade com PostgreSQL."),
    ], [1.35, 5.05])
    doc.add_paragraph("Os endpoints de criação cobrem contas, categorias, transações e orçamentos. Não há endpoint de pagamento ou de atualização de conta nesta versão.")

    doc.add_page_break()
    heading(doc, "2. Arquitetura e tecnologias", 1)
    doc.add_paragraph("A arquitetura escolhida é em camadas. Os módulos são agrupados por domínio para localizar regras e rotas relacionadas, enquanto a separação de apresentação, aplicação e infraestrutura reduz o acoplamento entre transporte HTTP, regras financeiras e persistência.")
    add_table(doc, ["Camada", "Responsabilidades", "Implementação"], [
        ("Apresentação", "Receber requisições, validar e serializar entradas e saídas.", "Controllers, DTOs, class-validator, Swagger"),
        ("Aplicação", "Executar casos de uso e aplicar regras do domínio.", "AuthService, AccountsService, TransactionsService e módulos financeiros"),
        ("Infraestrutura", "Persistir dados e manter estado operacional efêmero.", "TypeORM/PostgreSQL e ioredis/Redis"),
        ("Transversal", "Aplicar políticas que atravessam os domínios.", "JWT/RBAC, filtro RFC 7807, Correlation-ID e rate limiting"),
    ], [1.1, 2.6, 2.7])
    heading(doc, "Tecnologias", 2)
    add_table(doc, ["Tecnologia", "Uso no projeto"], [
        ("NestJS + TypeScript", "Framework HTTP, módulos e injeção de dependências."),
        ("TypeORM + PostgreSQL 16", "Mapeamento relacional e persistência financeira."),
        ("Redis 7 + ioredis", "Armazenamento de respostas e chaves de idempotência por 24 horas."),
        ("JWT + Passport + bcrypt", "Autenticação Bearer, estratégia JWT e hash de senha."),
        ("Swagger / OpenAPI", "Contrato navegável e execução de chamadas de teste."),
        ("Docker Compose", "Execução local da API e dos dois serviços de dados."),
    ], [2.0, 4.4])
    heading(doc, "Organização do código", 2)
    doc.add_paragraph("src/modules contém os módulos de autenticação, contas, categorias, transações, orçamentos e saúde. src/common contém componentes transversais como idempotência, contexto de requisição, autorização por papéis e o filtro global de erros.")

    doc.add_page_break()
    heading(doc, "3. C4 — Diagrama de Contêineres", 1)
    doc.add_paragraph("O diagrama mostra os elementos executáveis e os limites de responsabilidade: cliente, API NestJS, PostgreSQL e Redis. Swagger UI é servido pela própria API em /docs. Em desenvolvimento local a conexão HTTP pode não usar TLS; em produção, TLS deve ser terminado no ingresso ou balanceador.")
    doc.add_paragraph("Como ler: cada caixa representa um programa ou serviço que participa do sistema; as setas mostram para onde as informações são enviadas. O PostgreSQL é o arquivo organizado dos registros financeiros. O Redis é um apoio rápido para reconhecer pedidos repetidos.")
    add_image(doc, c4c, 6.55)
    caption(doc, "Figura 1. Contêineres e protocolos de comunicação do sistema.")
    heading(doc, "Justificativa dos bancos", 2)
    add_bullet(doc, "PostgreSQL guarda dados relacionais com integridade referencial: usuários, contas, categorias, transações e orçamentos. Valores monetários usam NUMERIC(14,2).")
    add_bullet(doc, "Redis mantém respostas idempotentes com expiração e reduz a necessidade de consultar o banco para repetir uma criação já concluída. Não é a fonte dos registros financeiros.")
    add_bullet(doc, "Docker Compose coordena serviços, health checks e volumes persistentes para desenvolvimento e demonstração local.")

    doc.add_page_break()
    heading(doc, "4. C4 — Diagrama de Componentes", 1)
    doc.add_paragraph("Dentro do contêiner NestJS, a apresentação direciona chamadas para os módulos de autenticação e financeiro. Componentes transversais aplicam políticas de erro, correlação, idempotência e vazão. Os módulos financeiros usam TypeORM; categorias e orçamentos acessam repositórios diretamente nos controllers nesta implementação.")
    doc.add_paragraph("Em linguagem simples, este desenho abre a caixa da API e mostra suas partes internas. Os componentes da esquerda recebem os pedidos, os do centro aplicam as regras e os da direita conversam com os bancos. A seta dos serviços transversais representa especificamente o uso do Redis pela idempotência; logs são enviados à saída do container e o limitador atual mantém contadores por processo.")
    add_image(doc, c4p, 6.55)
    caption(doc, "Figura 2. Componentes principais e dependências internas da API.")
    heading(doc, "Fluxo de uma requisição protegida", 2)
    for item in [
        "O middleware aceita ou gera um Correlation-ID e registra o resultado HTTP como JSON.",
        "Passport valida o JWT; o papel e o identificador do usuário autenticado ficam disponíveis à rota.",
        "O controller valida o DTO e encaminha o caso de uso ou consulta ao repositório.",
        "Criações críticas consultam Redis antes de executar e armazenam a resposta vinculada à chave e ao corpo da requisição.",
        "O filtro global converte exceções HTTP em application/problem+json.",
    ]: add_bullet(doc, item)

    doc.add_page_break()
    heading(doc, "5. Modelo de dados UML", 1)
    doc.add_paragraph("O modelo descreve as entidades e associações implementadas nos arquivos TypeORM. Um usuário possui seus próprios registros; cada transação pertence a uma conta e pode ter uma categoria; um orçamento pode limitar uma categoria específica ou ser geral.")
    doc.add_paragraph("No UML, cada retângulo representa um tipo de registro e lista algumas informações guardadas. As ligações representam relações: por exemplo, uma conta pode ter várias transações; uma transação pertence a uma conta e pode ter uma categoria. A versão textual abaixo do desenho explica as quantidades permitidas.")
    add_image(doc, uml, 6.55)
    caption(doc, "Figura 3. Diagrama UML de classes do domínio financeiro.")
    heading(doc, "Restrições e integridade", 2)
    add_bullet(doc, "E-mail de usuário é único; conta, categoria e orçamento têm índices compostos para evitar duplicações nos escopos definidos.")
    add_bullet(doc, "Transação exige uma conta pertencente ao mesmo usuário. Categoria, quando informada, também deve pertencer ao usuário e corresponder ao tipo receita/despesa.")
    add_bullet(doc, "A conta não pode ser removida enquanto houver transações relacionadas (integridade RESTRICT). A exclusão de usuário remove seus dados dependentes conforme cascatas definidas.")
    add_bullet(doc, "Valores persistidos como decimal no PostgreSQL são representados como strings pelos campos numéricos do TypeORM; entradas HTTP são validadas com até duas casas decimais.")

    doc.add_page_break()
    heading(doc, "6. Contrato da API", 1)
    doc.add_paragraph("Base local: http://localhost:3000/api/v1. A documentação OpenAPI é publicada em http://localhost:3000/docs. O prefixo e as credenciais podem ser configurados por variáveis de ambiente.")
    add_table(doc, ["Método e rota", "Acesso", "Descrição"], [
        ("POST /auth/register", "Público", "Registrar cliente; senha mínima de 8 caracteres."),
        ("POST /auth/login", "Público", "Validar credenciais e devolver token de acesso JWT."),
        ("GET /auth/admin/users", "Admin", "Listar usuários sem retornar hash de senha."),
        ("GET /health", "Público", "Verificar API e PostgreSQL."),
        ("GET/POST /accounts", "JWT", "Listar e criar contas; criação requer Idempotency-Key."),
        ("DELETE /accounts/{id}", "JWT", "Remover conta do usuário autenticado."),
        ("GET/POST /categories", "JWT", "Listar e criar categorias; criação requer Idempotency-Key."),
        ("GET/POST /transactions", "JWT", "Listar com filtros from/to e registrar movimentação idempotente."),
        ("DELETE /transactions/{id}", "JWT", "Remover transação do usuário autenticado."),
        ("GET/POST /budgets", "JWT", "Listar e criar orçamento mensal idempotente."),
    ], [2.15, .8, 3.45])
    heading(doc, "Exemplo de transação", 2)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.2)
    run = p.add_run('{\n  "kind": "expense",\n  "amount": 45.90,\n  "description": "Almoço",\n  "accountId": "UUID da conta",\n  "categoryId": "UUID da categoria (opcional)",\n  "occurredAt": "2026-09-27",\n  "status": "completed"\n}')
    run.font.name = "Consolas"; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph("As datas de occurredAt usam formato ISO AAAA-MM-DD. Os filtros from e to aceitam o mesmo formato e incluem as datas informadas.")

    doc.add_page_break()
    heading(doc, "7. Requisitos não funcionais", 1)
    add_table(doc, ["Requisito", "Implementação", "Como demonstrar"], [
        ("Autenticação", "JWT Bearer; senhas com bcrypt.", "Cadastro/login; rota protegida sem token deve responder 401."),
        ("Autorização", "Papéis admin e cliente; usuários isolados por ID.", "Cliente acessa seus dados; cliente em /auth/admin/users recebe 403."),
        ("Idempotência", "Header Idempotency-Key obrigatório em POST de contas, categorias, transações e orçamentos. Redis guarda a resposta por 24 horas.", "Mesma chave/corpo reproduz a resposta; mesma chave/corpo diferente retorna 409."),
        ("Erros", "Filtro global no formato RFC 7807 application/problem+json.", "Enviar DTO inválido e verificar status, title, detail, instance e requestId."),
        ("Observabilidade", "Correlation-ID/X-Request-ID propagado na resposta e log JSON por requisição.", "Enviar Correlation-ID em GET /health e localizar o valor no cabeçalho e logs."),
        ("Rate limiting", "5 requisições por minuto em login/registro; limite global de 100/minuto.", "Exceder o limite produz HTTP 429. O armazenamento atual é por processo."),
        ("OpenAPI", "Swagger UI em /docs com autenticação Bearer e parâmetros documentados.", "Executar operações no navegador Swagger."),
    ], [1.1, 3.0, 2.3])
    heading(doc, "Formato de erro", 2)
    p = doc.add_paragraph()
    run = p.add_run('{\n  "type": "https://httpstatuses.com/400",\n  "title": "Bad Request",\n  "status": 400,\n  "detail": "...",\n  "instance": "/api/v1/...",\n  "requestId": "..."\n}')
    run.font.name = "Consolas"; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor.from_string(NAVY)

    doc.add_page_break()
    heading(doc, "8. Execução e configuração", 1)
    heading(doc, "Inicialização local", 2)
    for item in [
        "Copiar .env.example para .env e substituir os segredos de exemplo.",
        "Executar docker compose up --build na raiz do projeto.",
        "Acompanhar docker compose ps; PostgreSQL e Redis devem ficar healthy e a API deve iniciar após ambos.",
        "Abrir Swagger em http://localhost:3000/docs e health check em http://localhost:3000/api/v1/health.",
    ]: add_bullet(doc, item)
    heading(doc, "Variáveis e dados", 2)
    add_table(doc, ["Variável", "Finalidade"], [
        ("JWT_SECRET", "Assinar tokens de acesso; usar segredo forte fora do ambiente local."),
        ("ADMIN_EMAIL / ADMIN_PASSWORD", "Provisionar o administrador inicial; senha com pelo menos 12 caracteres."),
        ("DB_HOST / DB_DATABASE / DB_USERNAME / DB_PASSWORD", "Configurar PostgreSQL. Dentro do Compose, o host do serviço é postgres."),
        ("REDIS_HOST / REDIS_PASSWORD", "Configurar Redis. Dentro do Compose, o host do serviço é redis."),
        ("DB_SYNCHRONIZE", "true facilita a demonstração local; produção deve usar false e migrações versionadas."),
    ], [2.15, 4.25])
    heading(doc, "Limites conhecidos e próximos passos", 2)
    add_bullet(doc, "Não há fluxo de refresh token implementado nesta versão; um token de acesso expirado requer novo login.")
    add_bullet(doc, "O rate limiter usa o armazenamento padrão por processo. Para várias réplicas, configurar armazenamento compartilhado Redis ou gateway.")
    add_bullet(doc, "Antes de produção: criar e executar migrações, configurar backup e restauração, segredos em cofre e TLS no ingresso.")
    add_bullet(doc, "As rotas oferecem registro de transações financeiras, mas ainda não implementam pagamentos, transferências ou atualização de conta/orçamento.")
    heading(doc, "Glossário", 2)
    add_table(doc, ["Termo", "Explicação sem jargão"], [
        ("API", "Serviço que recebe pedidos de um aplicativo e devolve respostas."),
        ("Endpoint / rota", "Endereço e ação disponíveis, como consultar contas ou registrar uma despesa."),
        ("Token JWT", "Credencial temporária emitida no login; funciona como um crachá para acessar rotas protegidas."),
        ("Idempotency-Key", "Identificador de uma tentativa de criação. Repetir a mesma tentativa com os mesmos dados não cria uma cópia."),
        ("PostgreSQL", "Banco de dados organizado em tabelas, usado como registro permanente dos dados financeiros."),
        ("Redis", "Banco rápido usado aqui para guardar temporariamente o resultado das tentativas idempotentes."),
        ("C4", "Modelo de diagramas que apresenta o sistema em níveis: serviços que o compõem e partes internas desses serviços."),
        ("UML", "Notação visual para mostrar os tipos de dados do projeto e como se relacionam."),
        ("DTO", "Formato esperado para os campos enviados em uma requisição; a API valida se os dados estão nesse formato."),
        ("TypeORM", "Ferramenta que traduz entre objetos do código e tabelas do PostgreSQL."),
    ], [1.6, 4.8])

    dest = OUT / "Documentacao_API_Gestao_Financeira_Apresentacao.docx"
    doc.save(dest)
    print(dest)


if __name__ == "__main__":
    build()
